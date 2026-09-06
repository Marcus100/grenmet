"""Create the September 2026 PRIC AST tide table as a Word document."""

import csv
from collections import defaultdict
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
PACKAGE = ROOT / "exports/02-atmosphere-aviation/meteorology-and-climate/tides/pric/2026-august-september-v1.2-pric-primary"
CSV_PATH = PACKAGE / "pric-tides-2026-08-01-to-2026-09-30.csv"
OUTPUT = PACKAGE / "pric-tides-september-2026-AST.docx"


def shade(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), fill)
    properties.append(element)


def set_cell_margins(cell, top=70, start=90, bottom=70, end=90) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    element = OxmlElement("w:tblHeader")
    element.set(qn("w:val"), "true")
    properties.append(element)


def set_fixed_width(table, widths) -> None:
    table.autofit = False
    table.allow_autofit = False
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)
            properties = cell._tc.get_or_add_tcPr()
            cell_width = properties.first_child_found_in("w:tcW")
            if cell_width is None:
                cell_width = OxmlElement("w:tcW")
                properties.append(cell_width)
            cell_width.set(qn("w:w"), str(round(width * 1440)))
            cell_width.set(qn("w:type"), "dxa")


def add_text(cell, text: str, *, bold=False, color=None, size=8.5) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def main() -> None:
    grouped: dict[str, list[str]] = defaultdict(list)
    with CSV_PATH.open(newline="", encoding="utf-8") as stream:
        for record in csv.DictReader(stream):
            if record["date_ast"][5:7] != "09":
                continue
            time_ast = record["time_ast"].removesuffix(" AST").split(" ")[-1]
            height = float(record["height_m_relative_uhslc_mllw_provisional_crosswalk"])
            grouped[record["date_ast"]].append(
                f"{record['event_type'].title()}\n{time_ast}\n{height:.2f} m"
            )

    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    for margin in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(section, margin, Inches(0.55))
    section.header_distance = Inches(0.25)
    section.footer_distance = Inches(0.25)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(2)
    run = title.add_run("Prickly Bay, Grenada")
    run.font.name = "Arial"
    run.font.size = Pt(18)
    run.bold = True
    run.font.color.rgb = RGBColor(31, 77, 120)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(6)
    run = subtitle.add_run("September 2026 Tide Table — Atlantic Standard Time (AST)")
    run.font.name = "Arial"
    run.font.size = Pt(11)
    run.bold = True

    note = document.add_paragraph()
    note.paragraph_format.space_after = Pt(6)
    note.paragraph_format.line_spacing = 1.0
    run = note.add_run(
        "INTERNAL GMS TECHNICAL-REVIEW CANDIDATE — NOT FOR NAVIGATION. "
        "Times are astronomical predictions from QC-clean PRIC radar observations. "
        "Heights are provisional UHSLC MLLW and require GMS benchmark/levelling approval."
    )
    run.font.name = "Arial"
    run.font.size = Pt(8.5)
    run.bold = True
    run.font.color.rgb = RGBColor(122, 90, 0)

    table = document.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    headers = ["Date", "Tide 1", "Tide 2", "Tide 3", "Tide 4"]
    header = table.rows[0]
    set_repeat_table_header(header)
    for cell, label in zip(header.cells, headers):
        shade(cell, "1F4D78")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        add_text(cell, label, bold=True, color="FFFFFF", size=9)

    for day in range(1, 31):
        row = table.add_row()
        date_key = f"2026-09-{day:02d}"
        events = grouped.get(date_key, [])[:4]
        events.extend(["—"] * (4 - len(events)))
        values = [str(day), *events]
        for cell, value in zip(row.cells, values):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, top=55, bottom=55)
            add_text(cell, value, bold=(cell is row.cells[0]), size=8.5)

    set_fixed_width(table, [0.55, 1.4875, 1.4875, 1.4875, 1.4875])

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("PRIC-primary v1.2 | Provisional internal GMS product | Generated 30 August 2026")
    footer_run.font.name = "Arial"
    footer_run.font.size = Pt(7.5)
    footer_run.font.color.rgb = RGBColor(85, 85, 85)

    document.core_properties.title = "Prickly Bay September 2026 Tide Table"
    document.core_properties.subject = "PRIC-primary astronomical tide predictions in AST"
    document.core_properties.author = "Grenada Meteorological Service"
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
